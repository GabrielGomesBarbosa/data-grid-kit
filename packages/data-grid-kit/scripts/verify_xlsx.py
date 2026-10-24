"""Opens the workbooks written by xlsx-fixture.ts with openpyxl (an independent implementation) and checks them."""
import sys
from openpyxl import load_workbook

d = sys.argv[1]
for kind in ("deflated", "stored"):
    wb = load_workbook(f"{d}/fixture-{kind}.xlsx")
    ws = wb.active
    assert "/" not in ws.title and "?" not in ws.title and len(ws.title) <= 31, ws.title
    merged = {str(r) for r in ws.merged_cells.ranges}
    assert {"A1:B1", "C1:D1"} <= merged, merged  # column groups
    assert ws["A1"].value == "People" and ws["C1"].value == "Pay"
    assert [c.value for c in ws[2]][:5] == ["Name", "Team", "Salary", "Hired", "Remote"]
    assert ws.freeze_panes == "A3", ws.freeze_panes
    assert ws["A3"].font.b and ws["A3"].value == "Data"
    assert ws.row_dimensions[4].outlineLevel == 1 and ws.row_dimensions[3].outlineLevel == 0
    assert ws["C4"].value == 1200.25 and ws["C4"].number_format == "$#,##0.00"
    assert ws["E4"].value is True and ws["E5"].value is False
    assert ws["D4"].is_date and ws["D4"].value.year == 2020 and ws["D4"].value.hour == 13
    assert ws["A4"].value == 'Ana <b> & "Co"'
    assert ws["A5"].value == "=1+1" and ws["A5"].data_type == "s", "text must never become a formula"
    assert ws["B5"].value == "@SUM(A1)" and ws["B5"].data_type == "s"
    assert ws["A6"].value == "Emoji 😀 日本語" and ws["B6"].value == "Data\nnew line"
    assert ws["C6"].value is None and ws["D6"].value is None and ws["E6"].value is None  # NaN, invalid date, undefined
    assert len(ws["A7"].value) == 32767, len(ws["A7"].value)  # cut at Excel's limit
    print(kind, "ok")

wb = load_workbook(f"{d}/fixture-features.xlsx")
assert wb.sheetnames == ["Items", "items (2)"], wb.sheetnames
ws = wb["Items"]
assert ws["A1"].fill.fgColor.rgb.endswith("112233") and ws["A1"].font.color.rgb.endswith("FFFFFF")
assert ws["A2"].font.i and ws["A2"].font.color.rgb.endswith("CC0000") and ws["A2"].alignment.horizontal == "center"
assert ws["C2"].font.b and ws["C2"].fill.fgColor.rgb.endswith("FFEE00")
assert ws["A3"].value == "=not a formula" and ws["A3"].data_type == "s", "text stays text"
assert ws["C4"].value == "=SUM(C2:C3)" and ws["C4"].data_type == "f", "an explicit formula is a formula"
dv = ws.data_validations.dataValidation
assert len(dv) == 1 and dv[0].type == "list" and dv[0].formula1 == '"open,done"' and str(dv[0].sqref) == "B2:B4"
assert wb["items (2)"]["A2"].value == "x"
print("features ok")
